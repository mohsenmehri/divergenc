"""Reproduce the four exact-header unique lists from the reviewed v5 workbook.
Requires openpyxl. Never opens/saves the source in write mode; no VBA executes.
Reviewed side-table bounds intentionally belong to this workbook snapshot.
"""
from pathlib import Path
from collections import OrderedDict
import hashlib
import json
import re
import unicodedata
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'Mohsen_FINAL_v5.xlsm'
OUTPUT = ROOT / 'exports/Mohsen_FINAL_v5_unique_fields.xlsx'
FIELDS = ['سرویس', 'نوع سرویس', 'پهنای باند', 'تجهیزات']
EXPECTED_SHA256 = 'bf1d5519e8e47b79212ef7bfbb7f0a6585faf6ae8604fd1f73a3af2a3b89ed66'
META = {'CONTROL_PANEL', 'CHANGE_LOG', 'SYSTEM_KEYS', 'SYSTEM_SHEET_CONFIG',
        'SYSTEM_DATA_DICTIONARY', 'SYSTEM_SEARCH_INDEX', 'SYSTEM_UNDO_BUFFER'}
# These are independent tables alongside other tables, not entire-sheet columns.
BOUNDS = {
    ('شبکه VSAT', 'G1'): 109, ('شبکه VSAT', 'O1'): 7,
    ('شبکه VSAT', 'G128'): 131,
    ('شبکه مبین نت 2', 'W49'): 59,
    ('شبکه آسیاتک2', 'T16'): 20,
    ('گزارش آماری شبکهVPN', 'C16'): 26,
    ('گزارش آماری شبکهVPN', 'C42'): 46,
    ('گزارش آماری شبکهVPN', 'C67'): 70,
    ('گزارش آماری شبکهVPN', 'C86'): 99,
}


def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def key(value):
    # Conservative text deduplication, not semantic service/unit conversion.
    s = unicodedata.normalize('NFC', text(value)).translate(str.maketrans('يك', 'یک'))
    return re.sub(r'\s+', ' ', s).strip()


def table_end(sheet, header):
    if (sheet.title, header.coordinate) in BOUNDS:
        return BOUNDS[sheet.title, header.coordinate], 'reviewed side-table boundary'
    # Stop before a new table/title. Blank rows can occur inside a table
    # (including merged records), so they must NOT terminate extraction. Never collect
    # values indefinitely down a column after it has become a different table.
    titles = {m.min_row for m in sheet.merged_cells.ranges
              if m.min_row == m.max_row and m.max_col - m.min_col >= 4
              and m.min_col <= header.column <= m.max_col}
    for r in range(header.row + 1, sheet.max_row + 1):
        values = [c.value for c in sheet[r]]
        if key(sheet.cell(r, header.column).value) == 'IP Free':
            return r - 1, 'new IP Free table in the same column'
        if key(values[0]) == 'ردیف' or r in titles:
            return r - 1, 'next header or merged section title'
    return sheet.max_row, 'end of worksheet'


def extract(source=SOURCE):
    workbook = load_workbook(source, data_only=True, keep_links=False)
    formulas = load_workbook(source, data_only=False, keep_links=False)
    result = {name: OrderedDict() for name in FIELDS}
    audit = []
    seen_cells = {name: set() for name in FIELDS}
    missing_caches = []
    for sheet in workbook:
        record = {'sheet': sheet.title, 'state': sheet.sheet_state, 'headers': [],
                  'role': 'metadata / controls (not data headers)' if sheet.title in META else 'data / report'}
        for row in sheet:
            for header in row:
                field = key(header.value)
                if field not in FIELDS:
                    continue
                if sheet.title in META:
                    record['headers'].append({'cell': header.coordinate, 'text': field,
                                              'excluded': 'field-name reference, not table header'})
                    continue
                end, reason = table_end(sheet, header)
                record['headers'].append({'cell': header.coordinate, 'text': field,
                                          'last_row': end, 'boundary': reason})
                for r in range(header.row + 1, end + 1):
                    cell = sheet.cell(r, header.column)
                    formula = formulas[sheet.title][cell.coordinate]
                    if formula.data_type == 'f' and cell.value is None:
                        missing_caches.append(f'{sheet.title}!{cell.coordinate}')
                    value = key(cell.value)
                    if not value:
                        continue
                    location = (sheet.title, cell.coordinate)
                    if location in seen_cells[field]:
                        continue
                    seen_cells[field].add(location)
                    item = result[field].setdefault(value, {'value': text(cell.value), 'sources': [], 'cells': [], 'variants': []})
                    if sheet.title not in item['sources']:
                        item['sources'].append(sheet.title)
                    if text(cell.value) not in item['variants']:
                        item['variants'].append(text(cell.value))
                    item['cells'].append(f'{sheet.title}!{cell.coordinate}')
        audit.append(record)
    assert not missing_caches, f'Uncached source formulas: {missing_caches}'
    return result, audit


def export(result):
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = 'مقادیر یکتای چهار ستون — Mohsen_FINAL_v5'
    wb.properties.description = 'Exact full headers only; duplicates removed within each field. Source workbook unchanged.'
    for index, (field, values) in enumerate(result.items(), 1):
        ws = wb.create_sheet(field)
        ws.sheet_view.rightToLeft = True
        ws.freeze_panes = 'A2'
        ws.append([field + ' — مقدار یکتا', 'شیت‌های مبدأ', 'اولین سلول مبدأ', 'تعداد سلول‌های مبدأ'])
        for item in values.values():
            ws.append([item['value'], ' | '.join(item['sources']), item['cells'][0], len(item['cells'])])
            # Explicit text, including strings beginning with =, prevents formula injection.
            for cell in ws[ws.max_row][:3]:
                cell.data_type = 's'
                cell.number_format = '@'
            ws.cell(ws.max_row, 1).comment = Comment(
                'همه سلول‌های مبدأ:\n' + '\n'.join(item['cells'])[:28000]
                + '\n\nصورت‌های دارای فاصله/حروف معادل:\n' + '\n'.join(item['variants']), 'Extraction audit')
        for cell in ws[1]:
            cell.font = Font(name='Calibri', size=12, color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='143D59')
            cell.alignment = Alignment(horizontal='right', vertical='center')
        ws['A1'].comment = Comment('فقط عنوان کامل ستون مطابق نام این شیت استخراج شده است؛ ستون‌های با پسوند یا عنوان طولانی‌تر وارد نشده‌اند. مقادیر پس از یکسان‌سازی فاصله‌های اضافی و ی/ک عربی-فارسی یکتا شده‌اند. تبدیل واحد، اصلاح املای داده یا ادغام معنایی انجام نشده است. متن‌های نامعمول، اعداد و خط‌تیره‌های موجود در داده حفظ شده‌اند.', 'روش استخراج')
        ws.row_dimensions[1].height = 30
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Calibri', size=11)
                cell.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True)
            ws.row_dimensions[row[0].row].height = min(100, 30 + 13 * row[0].value.count('\n'))
        for col, width in [('A', 65), ('B', 60), ('C', 48), ('D', 24)]:
            ws.column_dimensions[col].width = width
        table = Table(displayName=f'UniqueField{index}', ref=f'A1:D{ws.max_row}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '1:1'
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    check = load_workbook(OUTPUT, data_only=False)
    assert check.sheetnames == FIELDS
    for field, items in result.items():
        observed = [key(row[0].value) for row in check[field].iter_rows(min_row=2)]
        assert len(observed) == len(set(observed)) == len(items)
        assert set(observed) == set(items)
        assert not any(c.data_type == 'f' for row in check[field] for c in row)


if __name__ == '__main__':
    before = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    assert before == EXPECTED_SHA256, 'Source changed: review table boundaries before re-running.'
    result, audit = extract()
    export(result)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == before
    report = {'source': SOURCE.name, 'sha256': before, 'sheets_scanned': len(audit),
              'counts': {f: len(v) for f, v in result.items()}, 'sheet_audit': audit,
              'values': result}
    (ROOT / '.cache/extraction').mkdir(parents=True, exist_ok=True)
    (ROOT / '.cache/extraction/final-audit.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps({'sheets_scanned': len(audit), 'unique_counts': report['counts'],
                      'output': str(OUTPUT), 'source_sha256': before}, ensure_ascii=False, indent=2))
