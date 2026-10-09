import unittest
from openpyxl import Workbook, load_workbook
from scripts.extract_unique_excel_fields import FIELDS, SOURCE, OUTPUT, key, table_end, extract

class ExtractionTests(unittest.TestCase):
    def test_exact_header_matching(self):
        for value in ['تجهیزات ارتباطی', 'تجهیزات فیبر', 'شرح سرویس', 'نوع سرویس /شرح سرویس',
                      'پهنای باند-', 'پهنای باند(Kbps)', 'پهنای باند جدید']:
            self.assertNotIn(key(value), FIELDS)
        self.assertEqual(key('  نوع  سرويس '), 'نوع سرویس')

    def test_dedupe_is_conservative(self):
        self.assertEqual(key('  نوع\n سرويس  '), key('نوع سرویس'))
        self.assertEqual(key(128), key('128'))
        self.assertNotEqual(key('50 MB'), key('50 Mb'))
        self.assertNotEqual(key('2 Mb'), key('2048 Kb'))
        self.assertEqual(key(0), '0')

    def test_blank_rows_do_not_hide_later_records(self):
        ws=Workbook().active
        ws.append(['ردیف', 'سرویس'])
        ws.append([1,'MPLS'])
        ws.append([None,None])
        ws.append([2,'PTP'])
        self.assertEqual(table_end(ws,ws['B1'])[0],4)
        ws.append(['ردیف','تجهیزات'])
        ws.append([3,'Modem'])
        self.assertEqual(table_end(ws,ws['B1'])[0],4)

    def test_other_table_in_same_column_is_not_service_data(self):
        ws=Workbook().active
        ws.append(['ردیف','سرویس']);ws.append([1,'MPLS'])
        ws.append([None,' IP Free']);ws.append([None,'10.0.0.0/24'])
        self.assertEqual(table_end(ws,ws['B1'])[0],2)

    def test_source_and_export_coverage(self):
        result,audit=extract()
        self.assertEqual(len(audit),70)
        self.assertEqual({f:len(v) for f,v in result.items()},
                         {'سرویس':102,'نوع سرویس':44,'پهنای باند':107,'تجهیزات':13})
        wb=load_workbook(OUTPUT,data_only=False)
        self.assertEqual(wb.sheetnames,FIELDS)
        for field,items in result.items():
            values=[key(row[0].value)for row in wb[field].iter_rows(min_row=2)]
            self.assertEqual(len(values),len(set(values)))
            self.assertEqual(set(values),set(items))
            self.assertFalse(any(c.data_type=='f' for row in wb[field] for c in row))
        self.assertIn('شعب ادغامی درمردادوشهریور 1405!H26',result['سرویس']['MPLSاستانی']['cells'])
        self.assertNotIn('IP Free',result['سرویس'])
        self.assertNotIn('IP Free',result['پهنای باند'])
        self.assertEqual(set(s for v in result['تجهیزات'].values() for s in v['sources']),
                         {'سازمان های حاکمیتی','شبکه نوری ملل وفام','شبکه SIP ملل وفام'})

if __name__=='__main__':unittest.main()
